"""One Excel workbook for the whole run: a sheet with all invoice lines, one with the invoices, and one per table."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

INVOICE_LINE_COLS = ['fileName', 'page', 'lineIndex', 'description', 'sku', 'quantity', 'unit', 'unitPrice', 'discount',
                     'lineTaxRate', 'lineTaxAmount', 'lineTotal', 'lineCheck', 'invoiceNumber', 'invoiceDate', 'dueDate',
                     'supplierName', 'customerName', 'currency', 'total']
INVOICE_COLS = ['fileName', 'invoiceNumber', 'invoiceDate', 'dueDate', 'poNumber', 'supplierName', 'supplierTaxId',
                'customerName', 'currency', 'subtotal', 'taxRate', 'taxAmount', 'total', 'linesSum', 'totalsCheck', 'iban', 'sourceUrl']
MAX_SHEETS = 200


_WS = None  # the sheet being written (WriteOnlyCell needs it)


def _safe(v):
    """Text that starts like a formula ("-", "+", "=", "@") is stored as plain text with Excel's hidden quote prefix:
    the cell shows "- 4.5" and not "'- 4.5", and nothing is ever run as a formula."""
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v)
    if s[:1] in ('=', '+', '-', '@'):
        cell = WriteOnlyCell(_WS, value=s)
        cell.data_type = 's'
        cell.quotePrefix = True
        return cell
    return s


def build_workbook(rows: list[dict], invoices: list[dict]) -> bytes | None:
    global _WS
    if not rows and not invoices:
        return None
    wb = Workbook(write_only=True)
    lines = [r for r in rows if r.get('rowType') == 'invoice-line']
    if lines:
        ws = _WS = wb.create_sheet('Invoice lines')
        ws.append(INVOICE_LINE_COLS)
        for r in lines:
            ws.append([_safe(r.get(c)) for c in INVOICE_LINE_COLS])
    if invoices:
        ws = _WS = wb.create_sheet('Invoices')
        ws.append(INVOICE_COLS)
        for inv in invoices:
            ws.append([_safe(inv.get(c)) for c in INVOICE_COLS])
    tables = {}
    for r in rows:
        if r.get('rowType') == 'table-row':
            tables.setdefault((r['fileName'], r['tableIndex']), []).append(r)
    used = set()
    for n, ((fname, ti), trows) in enumerate(tables.items()):
        if n >= MAX_SHEETS:
            break
        base = re.sub(r'[\[\]:*?/\\]', '', re.sub(r'\.pdf$', '', fname, flags=re.I))[:20] or 'table'
        name = f'{base} t{ti}'[:31]
        k = 2
        while name in used:
            name = f'{base[:16]} t{ti}-{k}'[:31]
            k += 1
        used.add(name)
        ws = _WS = wb.create_sheet(name)
        first = trows[0]
        if first.get('tableTitle'):
            ws.append([_safe(first['tableTitle'])])
        cols = list(first['data'].keys())
        ws.append(cols if first.get('hasHeader') else [f'(no header) {c}' for c in cols])
        for r in trows:
            ws.append([_safe(r['data'].get(c)) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
