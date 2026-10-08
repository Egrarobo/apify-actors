"""Tests on the public sample PDFs in samples/ and the fixtures in tests/fixtures/.

Run:  python -m pytest -q tests   (needs pdfplumber, openpyxl, apify, httpx)
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from src.excel import build_workbook
from src.extract import PdfError, process_pdf
from src.invoice import map_roles
from src.numbers import detect_decimal_comma, parse_date, parse_number
from src.sources import is_pdf, not_pdf_reason, to_direct_url

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / 'samples'
FIXTURES = Path(__file__).resolve().parent / 'fixtures'


def run(name: str, **kw) -> dict:
    folder = SAMPLES if (SAMPLES / name).exists() else FIXTURES
    return process_pdf((folder / name).read_bytes(), **kw)


# ---------- numbers and dates ----------

@pytest.mark.parametrize('raw, dc, expected', [
    ('1,234.50', False, 1234.5),
    ('1.234,50', True, 1234.5),
    ('$24.90', False, 24.9),
    ('EUR 3.800,90', True, 3800.9),
    ('(15.00)', False, -15.0),
    ('-4.0', False, -4.0),
    ('6,5', True, 6.5),
    ('1,234', False, 1234.0),
    ('12.500', True, 12500.0),
    ('13,490', False, 13490.0),
    ('1 234,56', True, 1234.56),
    ('450,00 €', True, 450.0),
])
def test_parse_number(raw, dc, expected):
    assert parse_number(raw, dc)[0] == pytest.approx(expected)


@pytest.mark.parametrize('raw', ['61.00 - 94.00', 'PAP-A4-80', '9/25/2026', 'Total', '', '12 pcs extra'])
def test_not_a_number(raw):
    assert parse_number(raw) is None


def test_percent_and_currency():
    assert parse_number('21%') == (21.0, None, True)
    assert parse_number('$11.25')[1] == 'USD'
    assert parse_number('3.800,90 EUR', True)[1] == 'EUR'


@pytest.mark.parametrize('text, prefer, dc, expected', [
    ('September 15, 2026', 'auto', False, '2026-09-15'),
    ('15 Sept. 2026', 'auto', False, '2026-09-15'),
    ('30.09.2026', 'auto', True, '2026-09-30'),
    ('2026-10-14', 'auto', False, '2026-10-14'),
    ('03/04/2026', 'auto', False, '2026-03-04'),
    ('03/04/2026', 'DMY', False, '2026-04-03'),
    ('25/09/2026', 'auto', False, '2026-09-25'),
    ('no date here', 'auto', False, None),
])
def test_parse_date(text, prefer, dc, expected):
    assert parse_date(text, prefer, dc) == expected


def test_decimal_comma_detection():
    assert detect_decimal_comma('Total 1.234,50 EUR, VAT 21% 259,25, date 30.09.2026') is True
    assert detect_decimal_comma('Total $1,234.50, tax 98.76, stocks 13,490 and 707.1, note 3,13') is False


def test_header_roles():
    roles = map_roles(['#', 'Description', 'SKU', 'Qty', 'Unit Price', 'Amount'])
    assert roles == {0: 'lineNumber', 1: 'description', 2: 'sku', 3: 'quantity', 4: 'unitPrice', 5: 'lineTotal'}
    roles = map_roles(['Description', 'Quantity', 'Unit', 'Unit price', 'VAT', 'Total'])
    assert roles == {0: 'description', 1: 'quantity', 2: 'unit', 3: 'unitPrice', 4: 'taxRate', 5: 'lineTotal'}
    roles = map_roles(['Denumire', 'Cantitate', 'Preț unitar', 'Valoare'])
    assert set(roles.values()) == {'description', 'quantity', 'unitPrice', 'lineTotal'}


# ---------- invoices ----------

def test_us_invoice_with_ruled_table():
    r = run('invoice-us-ruled.pdf')
    assert r['documentType'] == 'invoice'
    inv = r['invoice']
    assert inv['invoiceNumber'] == 'INV-2026-0142'
    assert inv['invoiceDate'] == '2026-09-15'
    assert inv['dueDate'] == '2026-10-15'
    assert inv['poNumber'] == 'PO-7781'
    assert inv['supplierName'] == 'Larkspur Sample Supply Co.'
    assert inv['customerName'] == 'Bluebird Sample Bakery LLC'
    assert inv['currency'] == 'USD'
    assert (inv['subtotal'], inv['taxRate'], inv['taxAmount'], inv['total']) == (664.17, 8, 53.13, 717.3)
    assert inv['totalsCheck'] == 'ok'
    lines = [x for x in r['rows'] if x['rowType'] == 'invoice-line']
    assert len(lines) == 5
    assert lines[0]['description'] == 'Copy paper A4, 80 g, box of 5 reams'
    assert (lines[0]['sku'], lines[0]['quantity'], lines[0]['unitPrice'], lines[0]['lineTotal']) == ('PAP-A4-80', 12, 24.9, 298.8)
    assert all(x['lineCheck'] == 'ok' for x in lines)
    assert lines[0]['invoiceNumber'] == 'INV-2026-0142'  # invoice fields repeated on every line
    assert r['pagesWithData'] == [1]


def test_eu_invoice_borderless_two_pages_decimal_comma():
    r = run('invoice-eu-2-pages.pdf')
    inv = r['invoice']
    assert r['decimalComma'] is True
    assert inv['invoiceNumber'] == '2026/0387'
    assert (inv['invoiceDate'], inv['dueDate']) == ('2026-09-30', '2026-10-14')
    assert inv['supplierName'] == 'Wrenfield Sample Web Studio B.V.'
    assert inv['supplierTaxId'] == 'NL000000000B01'
    assert inv['customerName'] == 'Harbor Lane Sample Clinic'
    assert inv['currency'] == 'EUR'
    assert (inv['subtotal'], inv['taxRate'], inv['taxAmount'], inv['total']) == (3800.9, 21, 798.19, 4599.09)
    assert inv['iban'] == 'NL00EXMP0000000000'
    assert inv['totalsCheck'] == 'ok' and inv['linesSum'] == 3800.9
    lines = r['rows']
    assert len(lines) == 14
    assert {x['page'] for x in lines} == {1, 2}
    content = next(x for x in lines if x['description'] == 'Content updates')
    assert (content['quantity'], content['unit'], content['unitPrice'], content['lineTaxRate'], content['lineTotal']) == (6.5, 'h', 42, 21, 273)
    big = next(x for x in lines if x['description'] == 'Newsletter template design')
    assert big['unitPrice'] == 1150 and big['lineTotal'] == 1150
    assert all(x['lineCheck'] == 'ok' for x in lines)
    assert r['pagesWithData'] == [1, 2]


def test_invoice_in_tables_mode_gives_table_rows():
    r = run('invoice-us-ruled.pdf', mode='tables')
    assert r['documentType'] == 'tables'
    items = [x for x in r['rows'] if x.get('tableMethod') == 'lines']
    assert len(items) == 5
    assert items[0]['data'] == {'#': 1, 'Description': 'Copy paper A4, 80 g, box of 5 reams', 'SKU': 'PAP-A4-80', 'Qty': 12,
                                'Unit Price': 24.9, 'Amount': 298.8}


def test_invoice_mode_on_a_price_report_gives_no_rows():
    r = run('usda-shell-eggs-2026-10-02.pdf', mode='invoice')
    assert r['rows'] == [] and r['pagesWithData'] == []


# ---------- public price tables ----------

def test_usda_egg_prices_borderless_tables():
    r = run('usda-shell-eggs-2026-10-02.pdf')
    assert r['documentType'] == 'tables'
    tables = {}
    for x in r['rows']:
        tables.setdefault(x['tableIndex'], []).append(x)
    assert len(tables) == 7 and len(r['rows']) == 21
    first = tables[1]
    assert first[0]['tableTitle'] == 'National Shell Eggs - Caged · Delivered Warehouse, White, Cents Per Dozen'
    assert first[0]['data'] == {'Class': 'Extra Large', 'Price Range': '61.00 - 94.00', 'Average Price': 81.88,
                                'Price Change': 1, 'Last Reported (9/25/2026)': 80.88}
    assert [x['data']['Class'] for x in first] == ['Extra Large', 'Large', 'Medium']
    assert r['pagesWithData'] == [1, 2]


def test_eia_ruled_report_numbers_with_thousands_commas():
    r = run('eia-petroleum-balance-2026-10-02.pdf')
    assert r['decimalComma'] is False
    t1 = [x for x in r['rows'] if x['tableIndex'] == 1]
    crude = t1[0]['data']
    assert crude['Petroleum Stocks (Million Barrels)'] == 'Crude Oil'
    assert crude['Current Week 10/2/26'] == 707.1
    total = next(x['data'] for x in t1 if str(x['data']['Petroleum Stocks (Million Barrels)']).startswith('Total Stocks'))
    assert total['Current Week 10/2/26'] == 1520.4
    lower48 = next(x['data'] for x in r['rows'] if str(x['data'].get('col_1', '')).startswith('(3) Lower 48'))
    assert lower48['col_2'] == 13490


def test_parse_numbers_off_keeps_text():
    r = run('usda-shell-eggs-2026-10-02.pdf', parse_numbers=False)
    assert r['rows'][0]['data']['Average Price'] == '81.88'


def test_max_pages():
    r = run('invoice-eu-2-pages.pdf', max_pages=1)
    assert r['pagesRead'] == 1 and r['pagesWithData'] == [1]
    assert len(r['rows']) == 9
    assert any('Only the first 1 of 2 pages' in w for w in r['warnings'])


# ---------- errors ----------

def test_encrypted_pdf():
    with pytest.raises(PdfError, match='password'):
        run('encrypted-secret123.pdf')
    r = run('encrypted-secret123.pdf', password='secret123')
    assert r['invoice']['invoiceNumber'] == 'INV-2026-0142'


def test_scanned_pdf_has_no_rows():
    r = run('scanned-image-only.pdf')
    assert r['rows'] == [] and r['scannedPages'] == [1]


def test_not_a_pdf():
    with pytest.raises(PdfError):
        process_pdf(b'%PDF-1.4 this is not really a pdf')
    assert not is_pdf(b'<html><body>hi</body></html>')
    assert 'web page' in not_pdf_reason(b'<!DOCTYPE html><html>', 'x', 'text/html')
    assert 'Office' in not_pdf_reason(b'PK\x03\x04....', 'a.xlsx', '')


# ---------- sources and Excel ----------

@pytest.mark.parametrize('url, expected', [
    ('https://drive.google.com/file/d/ABC_123/view?usp=sharing', 'https://drive.usercontent.google.com/download?id=ABC_123&export=download&confirm=t'),
    ('https://www.dropbox.com/s/xyz/file.pdf?dl=0', 'https://www.dropbox.com/s/xyz/file.pdf?dl=1'),
    ('https://github.com/a/b/blob/main/x.pdf', 'https://raw.githubusercontent.com/a/b/main/x.pdf'),
    ('https://example.com/x.pdf', 'https://example.com/x.pdf'),
])
def test_share_links(url, expected):
    assert to_direct_url(url) == expected


def test_excel_workbook():
    inv = run('invoice-us-ruled.pdf')
    eggs = run('usda-shell-eggs-2026-10-02.pdf')
    rows = [{'fileName': 'invoice-us-ruled.pdf', **x} for x in inv['rows']] + \
           [{'fileName': 'usda-shell-eggs-2026-10-02.pdf', **x} for x in eggs['rows']]
    data = build_workbook(rows, [{'fileName': 'invoice-us-ruled.pdf', **inv['invoice']}])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames[:2] == ['Invoice lines', 'Invoices']
    assert len(wb.sheetnames) == 2 + 7
    lines = list(wb['Invoice lines'].iter_rows(values_only=True))
    assert len(lines) == 6 and lines[1][3] == 'Copy paper A4, 80 g, box of 5 reams'
    t1 = list(wb[wb.sheetnames[2]].iter_rows(values_only=True))
    assert t1[0][0].startswith('National Shell Eggs') and t1[1][0] == 'Class' and t1[2][2] == 81.88


# ---------- labels in other languages (synthetic lines, no PDF needed) ----------

def _lines(rows):
    """rows: list of (y, [(x, text), ...]) -> positioned lines, 10 pt font."""
    from src.tables import Word, build_lines
    words = []
    for y, parts in rows:
        for x, text in parts:
            x0 = x
            for w in text.split(' '):
                words.append(Word(x0, x0 + 5.5 * len(w), y, y + 10, w, 10.0, False))
                x0 += 5.5 * len(w) + 2.8  # one space
    return build_lines(words)


@pytest.mark.parametrize('rows, expected', [
    ([(50, [(40, 'Factura nr.: FX-2026-77')]), (65, [(40, 'Data emiterii: 05.10.2026')]), (80, [(40, 'Data scadenței: 20.10.2026')]),
      (95, [(40, 'Furnizor:'), (200, 'Exemplu Mobila SRL')]), (110, [(40, 'Cumpărător:'), (200, 'Client Exemplu SRL')])],
     {'invoiceNumber': 'FX-2026-77', 'invoiceDate': '2026-10-05', 'dueDate': '2026-10-20', 'supplierName': 'Exemplu Mobila SRL', 'customerName': 'Client Exemplu SRL'}),
    ([(50, [(40, 'Rechnungsnummer: RE-1009')]), (65, [(40, 'Rechnungsdatum: 01.10.2026')]), (80, [(40, 'Lieferant:'), (200, 'Beispiel GmbH')]),
      (95, [(40, 'Kunde:'), (200, 'Muster AG')])],
     {'invoiceNumber': 'RE-1009', 'invoiceDate': '2026-10-01', 'supplierName': 'Beispiel GmbH', 'customerName': 'Muster AG'}),
])
def test_labels_other_languages(rows, expected):
    from src.invoice import extract_invoice
    lines = _lines(rows)
    fields, items = extract_invoice([{'page': 1, 'lines': lines, 'height': 842.0, 'text': ''}], [], True)
    for k, v in expected.items():
        assert fields[k] == v, k
