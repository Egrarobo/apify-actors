"""Code columns stay text; text that starts like a formula is written without a visible apostrophe (S40, 10 Oct 2026)."""
import io

import openpyxl

from src.excel import build_workbook
from src.extract import CODE_COL_RE, _cell_value


def test_code_columns_stay_text():
    assert CODE_COL_RE.search('NAICS code') and CODE_COL_RE.search('SKU') and CODE_COL_RE.search('Item No')
    assert not CODE_COL_RE.search('Adjusted 2026 Aug. (a)')
    assert _cell_value('4411', False, True, code_column=True) == '4411'
    assert _cell_value('4411', False, True) == 4411
    assert _cell_value('784,846', False, True) == 784846


def test_excel_formula_like_text_has_no_visible_apostrophe():
    rows = [{'rowType': 'table-row', 'fileName': 'a.pdf', 'tableIndex': 1, 'hasHeader': True, 'tableTitle': None,
             'data': {'Kind': '- Gasoline stations', 'Value': 12.5, 'Note': '=SUM(A1)'}}]
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(rows, [])))
    ws = wb[wb.sheetnames[0]]
    values = [c.value for c in ws[2]]
    assert values == ['- Gasoline stations', 12.5, '=SUM(A1)']
    assert ws['A2'].data_type == 's' and ws['C2'].data_type == 's' and ws['C2'].quotePrefix
