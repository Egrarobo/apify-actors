"""Generate local test documents in test/fixtures (PDF, scanned PDF, DOCX, PPTX, XLSX, CSV, HTML, TXT)."""
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, ListFlowable, ListItem
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

import random
random.seed(7)
_WORDS = ("vector index embedding retrieval model prompt context window token budget latency cost quality "
          "chunk overlap heading table paragraph metadata source page citation answer grounding recall precision").split()


def sentence():
    w = random.sample(_WORDS, random.randint(6, 12))
    return w[0].capitalize() + ' ' + ' '.join(w[1:]) + '.'


def para(n=5):
    return ' '.join(sentence() for _ in range(n))


OUT = os.path.join(os.path.dirname(__file__), '..', 'test', 'fixtures')
os.makedirs(OUT, exist_ok=True)

FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
FONT_B = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
pdfmetrics.registerFont(TTFont('DV', FONT))
pdfmetrics.registerFont(TTFont('DVB', FONT_B))

LOREM = ("Retrieval-augmented generation (RAG) combines a search step with a language model. "
         "Documents are split into chunks, embedded, and stored in a vector database. "
         "At query time the most relevant chunks are retrieved and added to the prompt. ")

# ---------- PDF (text) ----------
ss = getSampleStyleSheet()
for k in ('Title', 'Heading1', 'Heading2', 'BodyText'):
    ss[k].fontName = 'DVB' if k != 'BodyText' else 'DV'
story = [Paragraph('Annual Report 2026', ss['Title']),
         Paragraph('1. Introduction', ss['Heading1']),
         Paragraph(LOREM * 3, ss['BodyText']),
         Paragraph('Diacritics test: Școala Robocode din Chișinău învață copiii să programeze. Кириллица тоже работает.', ss['BodyText']),
         Paragraph('1.1 Key numbers', ss['Heading2']),
         Table([['Quarter', 'Revenue', 'Students'], ['Q1', '12,500', '140'], ['Q2', '15,200', '171'], ['Q3', '18,900', '203']],
               style=TableStyle([('FONT', (0, 0), (-1, -1), 'DV'), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)])),
         Spacer(1, 12),
         ListFlowable([ListItem(Paragraph(t, ss['BodyText'])) for t in ('Robotics courses', 'Python for teens', 'Scratch for kids')], bulletType='bullet'),
         PageBreak(),
         Paragraph('2. Methodology', ss['Heading1']),
         Paragraph(para(8), ss['BodyText']), Paragraph(para(9), ss['BodyText']),
         Paragraph('2.1 Chunking strategy', ss['Heading2']),
         Paragraph(para(10), ss['BodyText']),
         PageBreak(),
         Paragraph('3. Conclusions', ss['Heading1']),
         Paragraph(para(6), ss['BodyText'])]


def footer(c, d):
    c.setFont('DV', 8)
    c.drawString(280, 20, f'Page {d.page}')


SimpleDocTemplate(os.path.join(OUT, 'report.pdf'), pagesize=A4, title='Annual Report 2026').build(story, onFirstPage=footer, onLaterPages=footer)

# ---------- Scanned PDF (image only) ----------
img = Image.new('RGB', (1654, 2339), 'white')  # A4 @ 200 dpi
d = ImageDraw.Draw(img)
fb = ImageFont.truetype(FONT_B, 64)
f = ImageFont.truetype(FONT, 40)
d.text((150, 200), 'Scanned Invoice No. 4471', font=fb, fill='black')
y = 380
for line in ['Customer: Robocode SRL, Chisinau', 'Service: Robotics course, 12 lessons', 'Total amount: 2,400 MDL',
             'Payment due within 14 days.', 'Thank you for your business!']:
    d.text((150, y), line, font=f, fill='black')
    y += 80
p = os.path.join(OUT, 'scanned.png')
img.save(p)
from reportlab.pdfgen import canvas
c = canvas.Canvas(os.path.join(OUT, 'scanned.pdf'), pagesize=A4)
c.drawImage(p, 0, 0, width=A4[0], height=A4[1])
c.showPage()
c.save()
os.remove(p)

# ---------- DOCX ----------
doc = Document()
doc.add_heading('Employee Handbook', 0)
doc.add_heading('Working hours', level=1)
doc.add_paragraph('Our office is open Monday to Friday, 9:00 to 18:00. ' + LOREM)
doc.add_heading('Benefits', level=2)
for t in ('Health insurance', 'Training budget of 500 EUR per year', 'Flexible remote work'):
    doc.add_paragraph(t, style='List Bullet')
p = doc.add_paragraph('Important: ')
p.add_run('bold text').bold = True
p.add_run(' and ')
p.add_run('italic text').italic = True
doc.add_heading('Holiday table', level=2)
t = doc.add_table(rows=3, cols=3)
for r, row in enumerate([['Holiday', 'Date', 'Paid'], ['New Year', 'Jan 1', 'Yes'], ['Independence Day', 'Aug 27', 'Yes']]):
    for ci, v in enumerate(row):
        t.cell(r, ci).text = v
doc.add_heading('Long section', level=1)
for _ in range(8):
    doc.add_paragraph(para(9))
doc.save(os.path.join(OUT, 'handbook.docx'))

# ---------- PPTX ----------
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Product Launch Plan'
s.placeholders[1].text = 'Q4 2026 go-to-market'
s.notes_slide.notes_text_frame.text = 'Speaker note: welcome everyone and introduce the team.'
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'Timeline'
tf = s.placeholders[1].text_frame
tf.text = 'October: beta testing'
for t in ('November: public launch', 'December: first 1,000 customers'):
    tf.add_paragraph().text = t
s.notes_slide.notes_text_frame.text = 'Emphasize the December milestone.'
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = 'Budget'
tbl = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
for r, row in enumerate([['Item', 'Cost'], ['Ads', '$5,000'], ['Events', '$3,000']]):
    for ci, v in enumerate(row):
        tbl.cell(r, ci).text = v
prs.save(os.path.join(OUT, 'launch.pptx'))

# ---------- XLSX ----------
wb = Workbook()
ws = wb.active
ws.title = 'Sales'
for row in [['Region', 'Q1', 'Q2', 'Total'], ['North', 100, 120, None], ['South', 90, 95, None], ['East | West', 50, 70, None]]:
    ws.append(row)
for r in range(2, 5):
    ws[f'D{r}'] = f'=B{r}+C{r}'
ws2 = wb.create_sheet('Staff')
ws2.append(['Name', 'Role', 'Start date'])
import datetime
ws2.append(['Ana Popescu', 'Teacher', datetime.date(2024, 9, 1)])
ws2.append(['Ion Rusu', 'Engineer', datetime.date(2025, 2, 15)])
wb.save(os.path.join(OUT, 'sales.xlsx'))

# ---------- CSV / HTML / TXT ----------
with open(os.path.join(OUT, 'products.csv'), 'w', encoding='utf-8') as fh:
    fh.write('sku,name,price\nA1,"Arduino kit, starter",45\nB2,Raspberry Pi 5,80\n')
with open(os.path.join(OUT, 'article.html'), 'w', encoding='utf-8') as fh:
    fh.write('''<!doctype html><html><head><title>How RAG works</title><style>body{}</style><script>var x=1;</script></head>
<body><nav><a href="/">Home</a> | <a href="/blog">Blog</a></nav>
<main><article><h1>How RAG works</h1><p>RAG stands for <strong>retrieval-augmented generation</strong>. See <a href="https://example.com/rag">this guide</a>.</p>
<h2>Steps</h2><ol><li>Split documents</li><li>Embed chunks</li><li>Retrieve and generate</li></ol>
<h2>Comparison</h2><table><thead><tr><th>Approach</th><th>Cost</th></tr></thead><tbody><tr><td>Fine-tuning</td><td>High</td></tr><tr><td>RAG</td><td>Low</td></tr></tbody></table>
<pre><code>chunks = split(doc, size=1000)</code></pre></article></main><footer>(c) 2026 Example Inc.</footer></body></html>''')
with open(os.path.join(OUT, 'notes.txt'), 'w', encoding='utf-8') as fh:
    fh.write('Meeting notes\n\nAttendees: Silvia, Dan\n\n' + LOREM * 2 + '\n')
print('fixtures written to', os.path.abspath(OUT))
